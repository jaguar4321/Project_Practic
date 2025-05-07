import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill, Color
from openpyxl.utils import get_column_letter
from .models import *
import os
from django.db.models import Count


# def create_excel_template(visits, selected_student=None):
#     file_name = 'excel_template.xlsx'
#     if os.path.exists(file_name):
#         os.remove(file_name)
#
#     wb = Workbook()
#     ws = wb.active
#     ws.title = "Відвідуваність"
#
#     # Стилі
#     bold_font = Font(bold=True)
#     thin_border = Border(left=Side(style='thin'),
#                          right=Side(style='thin'),
#                          top=Side(style='thin'),
#                          bottom=Side(style='thin'))
#     center_alignment = Alignment(horizontal='center', vertical='center')
#
#     # Додаємо заголовок таблиці (один раз)
#     ws.append(['№', 'Група', 'Студент'])
#     for cell in ws[1]:
#         cell.font = bold_font
#         cell.border = thin_border
#         cell.alignment = center_alignment
#
#     # Заморожуємо перші 3 стовпці та перший рядок
#     ws.freeze_panes = 'D2'
#
#     # Початкова ширина стовпців
#     ws.column_dimensions['A'].width = 5
#     ws.column_dimensions['B'].width = 10
#     ws.column_dimensions['C'].width = 15
#
#     current_row = 2
#
#     # Якщо немає відвідувань, створюємо порожню таблицю
#     if not visits.exists():
#         ws.cell(row=current_row, column=1, value="Немає даних про відвідування").font = bold_font
#         ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
#         max_length = len("Немає даних про відвідування")
#         adjusted_width = max_length + 5
#         total_width = ws.column_dimensions['A'].width + ws.column_dimensions['B'].width + ws.column_dimensions['C'].width
#         if adjusted_width > total_width:
#             extra_width = (adjusted_width - total_width) / 2
#             ws.column_dimensions['B'].width = max(extra_width, 10)
#             ws.column_dimensions['C'].width = max(extra_width, 15)
#         wb.save(file_name)
#         return os.path.abspath(file_name)
#
#     # Збираємо всі назви для визначення максимальної ширини об’єднаних комірок
#     max_merged_length = 0
#     longest_title = ""
#
#     # Отримуємо всі унікальні дисципліни
#     disciplines = visits.values('discipline').distinct()
#
#     for disc in disciplines:
#         discipline = Discipline.objects.get(id=disc['discipline'])
#
#         # Отримуємо всі унікальні види занять для цієї дисципліни
#         lessons = visits.filter(discipline=discipline).values('lesson').distinct()
#
#         for lesson in lessons:
#             lesson_type = lesson['lesson']
#
#             # Додаємо назву дисципліни та виду заняття
#             title = f"{discipline.name} - {lesson_type}"
#             ws.cell(row=current_row, column=1, value=title).font = bold_font
#             ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
#             if len(title) > max_merged_length:
#                 max_merged_length = len(title)
#                 longest_title = title
#             current_row += 1
#
#             # Отримуємо унікальні групи з відвідувань для цієї дисципліни та виду заняття
#             groups = visits.filter(discipline=discipline, lesson=lesson_type).values('group').distinct()
#
#             # Якщо обрано студентів, використовуємо їх
#             if selected_student is not None and selected_student.exists():
#                 all_students = selected_student
#             else:
#                 # Інакше беремо всіх студентів із груп, які є у відвідуваннях
#                 group_ids = [group['group'] for group in groups]
#                 all_students = Student.objects.filter(group__id__in=group_ids).order_by('group__name', 'full_name')
#
#             # Фільтруємо студентів, які мають відвідування для цієї комбінації "Дисципліна - Вид заняття"
#             students_with_visits = []
#             for student in all_students:
#                 if visits.filter(email=student, discipline=discipline, lesson=lesson_type).exists():
#                     students_with_visits.append(student)
#
#             if not students_with_visits:
#                 continue  # Пропускаємо, якщо немає студентів із відвідуваннями для цієї комбінації
#
#             # Групуємо студентів за їхньою групою
#             grouped_students = {}
#             for student in students_with_visits:
#                 group_name = student.group.name
#                 if group_name not in grouped_students:
#                     grouped_students[group_name] = []
#                 grouped_students[group_name].append(student)
#
#             # Перебираємо групи студентів
#             for group_name, group_students in sorted(grouped_students.items()):
#                 # Додаємо назву групи
#                 group_title = f"Група: {group_name}"
#                 ws.cell(row=current_row, column=1, value=group_title).font = bold_font
#                 ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
#                 if len(group_title) > max_merged_length:
#                     max_merged_length = len(group_title)
#                     longest_title = group_title
#                 current_row += 1
#
#                 # Отримуємо дати для цієї дисципліни та виду заняття
#                 dates = visits.filter(discipline=discipline, lesson=lesson_type).values('date').distinct().order_by('date')
#
#                 # Додаємо рядок із датами
#                 if dates:
#                     column = 4
#                     for date in dates:
#                         formatted_date = date['date'].strftime('%d/%m/%Y')
#                         ws.cell(row=current_row, column=column, value=formatted_date)
#                         ws.cell(row=current_row, column=column).font = bold_font
#                         ws.cell(row=current_row, column=column).border = thin_border
#                         ws.cell(row=current_row, column=column).alignment = center_alignment
#                         column += 1
#                     current_row += 1
#
#                 # Додаємо студентів та їх відвідуваність
#                 student_counter = 1
#                 for student in group_students:
#                     ws.cell(row=current_row, column=1, value=student_counter)
#                     ws.cell(row=current_row, column=2, value=student.group.name)
#                     ws.cell(row=current_row, column=3, value=student.full_name)
#
#                     # Стилі для комірок студента
#                     for col in range(1, 4):
#                         ws.cell(row=current_row, column=col).border = thin_border
#                         ws.cell(row=current_row, column=col).alignment = center_alignment
#
#                     # Заповнюємо відвідуваність
#                     column = 4
#                     for date in dates:
#                         visit_exists = visits.filter(
#                             email=student,
#                             discipline=discipline,
#                             lesson=lesson_type,
#                             date=date['date']
#                         ).exists()
#                         ws.cell(row=current_row, column=column, value='+' if visit_exists else '-')
#                         ws.cell(row=current_row, column=column).border = thin_border
#                         ws.cell(row=current_row, column=column).alignment = center_alignment
#                         column += 1
#
#                     student_counter += 1
#                     current_row += 1
#
#                 current_row += 1  # Порожній рядок між групами
#
#             current_row += 1  # Порожній рядок між видами занять
#
#         current_row += 1  # Порожній рядок між дисциплінами
#
#     # Налаштування ширини для об’єднаних комірок (стовпці A–C)
#     adjusted_width = max_merged_length + 5
#     total_width = ws.column_dimensions['A'].width + ws.column_dimensions['B'].width + ws.column_dimensions['C'].width
#     if adjusted_width > total_width:
#         extra_width = (adjusted_width - total_width) / 2
#         ws.column_dimensions['B'].width = max(extra_width, 10)
#         ws.column_dimensions['C'].width = max(extra_width, 15)
#
#     # Автоматична ширина решти стовпців (починаючи з D)
#     for col in range(4, ws.max_column + 1):
#         column_letter = get_column_letter(col)
#         max_length = 0
#         for row in range(1, ws.max_row + 1):
#             cell = ws.cell(row=row, column=col)
#             if cell.value and not isinstance(cell, openpyxl.cell.cell.MergedCell):
#                 try:
#                     max_length = max(max_length, len(str(cell.value)))
#                 except:
#                     pass
#         adjusted_width = (max_length + 2) * 1.2
#         ws.column_dimensions[column_letter].width = adjusted_width
#
#     wb.save(file_name)
#     return os.path.abspath(file_name)

def create_excel_template(visits, selected_student=None):
    file_name = 'excel_template.xlsx'
    if os.path.exists(file_name):
        os.remove(file_name)

    wb = Workbook()
    ws = wb.active
    ws.title = "Відвідуваність"

    # Стили
    bold_font = Font(bold=True)
    thin_border = Border(left=Side(style='thin'),
                         right=Side(style='thin'),
                         top=Side(style='thin'),
                         bottom=Side(style='thin'))
    center_alignment = Alignment(horizontal='center', vertical='center')

    # Добавляем заголовок таблицы (один раз)
    ws.append(['№', 'Група', 'Студент'])
    for cell in ws[1]:
        cell.font = bold_font
        cell.border = thin_border
        cell.alignment = center_alignment

    # Замораживаем первые 3 столбца и первый рядок
    ws.freeze_panes = 'D2'

    # Начальная ширина столбцов
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 15

    current_row = 2
    max_total_width = ws.column_dimensions['A'].width + ws.column_dimensions['B'].width + ws.column_dimensions['C'].width  # Отслеживаем максимальную ширину

    # Если нет посещений, создаем пустую таблицу
    if not visits.exists():
        ws.cell(row=current_row, column=1, value="Немає даних про відвідування").font = bold_font
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
        max_length = len("Немає даних про відвідування")
        adjusted_width = max_length + 5
        total_width = ws.column_dimensions['A'].width + ws.column_dimensions['B'].width + ws.column_dimensions['C'].width
        if adjusted_width > total_width:
            extra_width = (adjusted_width - total_width) / 2
            ws.column_dimensions['B'].width = extra_width
            ws.column_dimensions['C'].width = extra_width
        wb.save(file_name)
        return os.path.abspath(file_name)

    # Получаем все уникальные дисциплины
    disciplines = visits.values('discipline').distinct()

    for disc in disciplines:
        discipline = Discipline.objects.get(id=disc['discipline'])

        # Получаем все уникальные виды занятий для этой дисциплины
        lessons = visits.filter(discipline=discipline).values('lesson').distinct()

        for lesson in lessons:
            lesson_type = lesson['lesson']

            # Добавляем название дисциплины и вида занятия
            title = f"{discipline.name} - {lesson_type}"
            ws.cell(row=current_row, column=1, value=title).font = bold_font
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)

            # Настраиваем ширину для этой конкретной строки
            title_length = len(title)
            adjusted_width = title_length + 5
            if adjusted_width > max_total_width:
                max_total_width = adjusted_width
                extra_width = (adjusted_width - ws.column_dimensions['A'].width) / 2
                ws.column_dimensions['B'].width = extra_width
                ws.column_dimensions['C'].width = extra_width

            current_row += 1

            # Получаем уникальные группы с посещениями для этой дисциплины и вида занятия
            groups = visits.filter(discipline=discipline, lesson=lesson_type).values('group').distinct()

            # Если выбраны студенты, используем их
            if selected_student is not None and selected_student.exists():
                all_students = selected_student
            else:
                # Иначе берем всех студентов из групп, которые есть в посещениях
                group_ids = [group['group'] for group in groups]
                all_students = Student.objects.filter(group__id__in=group_ids).order_by('group__name', 'full_name')

            # Фильтруем студентов, которые имеют посещения для этой комбинации "Дисциплина - Вид занятия"
            students_with_visits = []
            for student in all_students:
                if visits.filter(email=student, discipline=discipline, lesson=lesson_type).exists():
                    students_with_visits.append(student)

            if not students_with_visits:
                continue  # Пропускаем, если нет студентов с посещениями для этой комбинации

            # Группируем студентов по их группе
            grouped_students = {}
            for student in students_with_visits:
                group_name = student.group.name
                if group_name not in grouped_students:
                    grouped_students[group_name] = []
                grouped_students[group_name].append(student)

            # Перебираем группы студентов
            for group_name, group_students in sorted(grouped_students.items()):
                # Добавляем название группы
                group_title = f"Група: {group_name}"
                ws.cell(row=current_row, column=1, value=group_title).font = bold_font
                ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)

                # Настраиваем ширину для этой конкретной строки группы
                group_title_length = len(group_title)
                adjusted_width = group_title_length + 5
                if adjusted_width > max_total_width:
                    max_total_width = adjusted_width
                    extra_width = (adjusted_width - ws.column_dimensions['A'].width) / 2
                    ws.column_dimensions['B'].width = extra_width
                    ws.column_dimensions['C'].width = extra_width

                current_row += 1

                # Получаем даты для этой дисциплины и вида занятия
                dates = visits.filter(discipline=discipline, lesson=lesson_type).values('date').distinct().order_by('date')

                # Добавляем рядок с датами
                if dates:
                    column = 4
                    for date in dates:
                        formatted_date = date['date'].strftime('%d/%m/%Y')
                        ws.cell(row=current_row, column=column, value=formatted_date)
                        ws.cell(row=current_row, column=column).font = bold_font
                        ws.cell(row=current_row, column=column).border = thin_border
                        ws.cell(row=current_row, column=column).alignment = center_alignment
                        column += 1
                    current_row += 1

                # Добавляем студентов и их посещаемость
                student_counter = 1
                for student in group_students:
                    ws.cell(row=current_row, column=1, value=student_counter)
                    ws.cell(row=current_row, column=2, value=student.group.name)
                    ws.cell(row=current_row, column=3, value=student.full_name)

                    # Стили для ячеек студента
                    for col in range(1, 4):
                        ws.cell(row=current_row, column=col).border = thin_border
                        ws.cell(row=current_row, column=col).alignment = center_alignment

                    # Заполняем посещаемость
                    column = 4
                    for date in dates:
                        visit_exists = visits.filter(
                            email=student,
                            discipline=discipline,
                            lesson=lesson_type,
                            date=date['date']
                        ).exists()
                        ws.cell(row=current_row, column=column, value='+' if visit_exists else '-')
                        ws.cell(row=current_row, column=column).border = thin_border
                        ws.cell(row=current_row, column=column).alignment = center_alignment
                        column += 1

                    student_counter += 1
                    current_row += 1

                current_row += 1  # Пустая строка между группами

            current_row += 1  # Пустая строка между видами занятий

        current_row += 1  # Пустая строка между дисциплинами

    # Автоматическая ширина остальных столбцов (начиная с D)
    for col in range(4, ws.max_column + 1):
        column_letter = get_column_letter(col)
        max_length = 0
        for row in range(1, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            if cell.value and not isinstance(cell, openpyxl.cell.cell.MergedCell):
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
        adjusted_width = (max_length + 2) * 1.2
        ws.column_dimensions[column_letter].width = adjusted_width

    wb.save(file_name)
    return os.path.abspath(file_name)